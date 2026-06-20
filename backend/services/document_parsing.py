"""Document ingestion: PDF (PyMuPDF) with OCR fallback, DOCX (python-docx), text.

Heavy libraries are imported lazily inside each parser so the module imports even
where a given backend isn't installed, and so plain-text ingestion has no native
dependencies.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Optional


@dataclass
class ParsedDocument:
    text: str
    source_format: str  # "pdf" | "docx" | "text"
    page_count: Optional[int] = None
    ocr_used: bool = False


def detect_language(text: str) -> str:
    """Lightweight EN/AR detection by Unicode block frequency."""
    arabic = 0
    latin = 0
    for ch in text:
        if "؀" <= ch <= "ۿ" or "ݐ" <= ch <= "ݿ" or "ﭐ" <= ch <= "﷿":
            arabic += 1
        elif ("a" <= ch <= "z") or ("A" <= ch <= "Z"):
            latin += 1
    return "ar" if arabic > latin else "en"


def parse_text(data: bytes | str) -> ParsedDocument:
    text = data.decode("utf-8", errors="ignore") if isinstance(data, bytes) else data
    return ParsedDocument(text=text.strip(), source_format="text")


def parse_docx(data: bytes) -> ParsedDocument:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    text = "\n".join(part for part in parts if part is not None)
    return ParsedDocument(text=text.strip(), source_format="docx")


def _ocr_page(page) -> str:
    try:
        import pytesseract
        from PIL import Image

        pix = page.get_pixmap(dpi=200)
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(image)
    except Exception:
        return ""


def parse_pdf(data: bytes, *, ocr: bool = True) -> ParsedDocument:
    import fitz  # PyMuPDF

    document = fitz.open(stream=data, filetype="pdf")
    try:
        page_texts: list[str] = []
        ocr_used = False
        for page in document:
            text = page.get_text("text")
            if (not text or len(text.strip()) < 20) and ocr:
                ocr_text = _ocr_page(page)
                if ocr_text.strip():
                    text = ocr_text
                    ocr_used = True
            page_texts.append(text)
        page_count = document.page_count
    finally:
        document.close()
    return ParsedDocument(
        text="\n\n".join(page_texts).strip(),
        source_format="pdf",
        page_count=page_count,
        ocr_used=ocr_used,
    )


def parse_document(filename: str, data: bytes, content_type: Optional[str] = None) -> ParsedDocument:
    name = (filename or "").lower()
    ctype = (content_type or "").lower()
    if name.endswith(".pdf") or "pdf" in ctype:
        return parse_pdf(data)
    if name.endswith(".docx") or "officedocument.wordprocessing" in ctype:
        return parse_docx(data)
    if name.endswith(".doc"):
        raise ValueError("Legacy .doc is not supported — please upload .docx or PDF.")
    return parse_text(data)
